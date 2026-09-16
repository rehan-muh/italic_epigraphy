#!/usr/bin/env python3
"""
Build abstract/qdacl/qdlca27-abstract.docx, the file the QDLCA27 call asks for.

The text is converted from abstract/qdacl/qdlca27-abstract.tex rather than
retyped, the figures are the PNGs scripts/make_qdlca_figures.py writes into
the same folder, and the reference list is lifted from the compiled PDF so
that biblatex, not this script, decides how a reference is formatted. Compile
the PDF first (pdflatex / bibtex / pdflatex x2).

Submission is anonymous, so the document properties are blanked: Word
otherwise writes the machine's registered user into author/last-modified-by.

    python scripts/make_qdlca_docx.py      (python-docx + pymupdf on py 3.14)
"""
from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path

import docx
import fitz
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DIR = ROOT / "abstract" / "qdacl"
TEX = DIR / "qdlca27-abstract.tex"
PDF = DIR / "qdlca27-abstract.pdf"
BIB = DIR / "references.bib"
OUT = DIR / "qdlca27-abstract.docx"

MATH = {r"\pm": "\u00b1", r"\times": "\u00d7", r"^{\circ}": "\u00b0",
        r"\,": "\u2009"}
ACCENT = {"`": "\u0300", "'": "\u0301", "^": "\u0302", '"': "\u0308", "~": "\u0303"}


# ------------------------------------------------------------ citations
def cite_labels():
    """bib key -> 'Author Year' as an authoryear style prints it."""
    src = BIB.read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"@\w+\{([^,]+),(.*?)\n\}", src, re.S):
        key, body = m.group(1).strip(), m.group(2)
        author = re.search(r"(?:author|editor)\s*=\s*\{(.*?)\}\s*,?\n", body, re.S).group(1)
        year = re.search(r"year\s*=\s*\{?(\d{4})", body).group(1)
        names = [a.strip() for a in re.split(r"\s+and\s+", author)]
        last = [n.split(",")[0].strip() if "," in n else n.split()[-1] for n in names]
        last = [re.sub(r"[{}]", "", l) for l in last]
        if len(last) == 1:
            who = last[0]
        elif len(last) == 2:
            who = f"{last[0]} and {last[1]}"
        else:
            who = f"{last[0]} et al."
        out[key] = f"{who} {year}"
    return out


# ---------------------------------------------------------------- detex
def detex(s: str, cites: dict, figs: dict) -> str:
    s = re.sub(r"\\parencite\{([^}]*)\}",
               lambda m: "(" + "; ".join(cites[k.strip()] for k in m.group(1).split(",")) + ")", s)
    s = re.sub(r"Fig\.~\\ref\{([^}]*)\}", lambda m: f"Fig.\u00a0{figs[m.group(1)]}", s)
    s = re.sub(r"\\ref\{([^}]*)\}", lambda m: str(figs[m.group(1)]), s)

    def math(m):
        t = m.group(1)
        for k, v in MATH.items():
            t = t.replace(k, v)
        return t
    s = re.sub(r"\$(.+?)\$", math, s)
    s = s.replace(r"\,", "\u2009")
    s = re.sub(r"\\([`'^\"~])\{?([a-zA-Z])\}?",
               lambda m: unicodedata.normalize("NFC", m.group(2) + ACCENT[m.group(1)]), s)
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("{,}", ",").replace(r"\%", "%").replace("~", "\u00a0")
    s = s.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", s).strip()


def runs(s: str, cites, figs):
    """(text, bold) runs; only \\textbf survives as markup in this abstract."""
    out, pos = [], 0
    for m in re.finditer(r"\\textbf\{([^{}]*)\}", s):
        if m.start() > pos:
            out.append((detex(s[pos:m.start()], cites, figs) + " ", False))
        out.append((detex(m.group(1), cites, figs), True))
        pos = m.end()
    if pos < len(s):
        out.append((" " + detex(s[pos:], cites, figs), False))
    return out


# ---------------------------------------------------------------- parse
def parse(tex: str):
    src = "\n".join(l for l in tex.splitlines() if not l.lstrip().startswith("%"))
    body_src = src.split(r"\begin{document}", 1)[1].split(r"\clearpage", 1)[0]
    title = re.search(r"\\begin\{center\}\s*\\textbf\{(.*?)\}\s*\\end\{center\}",
                      body_src, re.S).group(1).strip()
    body_src = re.sub(r"\\begin\{center\}.*?\\end\{center\}", "", body_src, flags=re.S)
    body_src = re.sub(r"\\vspace\{[^}]*\}", "", body_src)
    paras = [p.strip() for p in body_src.split("\n\n") if p.strip()]

    figs_src = src.split(r"\clearpage", 1)[1]
    figs = []
    for m in re.finditer(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}\s*\\caption\{(.*?)\}\s*\\label\{([^}]*)\}",
                         figs_src, re.S):
        figs.append((DIR / m.group(1), m.group(2), m.group(3)))
    return title, paras, figs


def references_from_pdf():
    text = "\n".join(p.get_text() for p in fitz.open(PDF))
    refs = text.split("References", 1)[1].strip()
    # one reference per entry: entries start with 'Surname, I.' at a line start
    entries, cur = [], []
    for line in refs.splitlines():
        if re.match(r"^[A-Z][^,]*, [A-Z]\.", line) and cur:
            entries.append(" ".join(cur))
            cur = []
        cur.append(line.strip())
    if cur:
        entries.append(" ".join(cur))
    return [re.sub(r"-\s+(?=[a-z])", "", e) for e in entries]


# ---------------------------------------------------------------- write
def add(doc, parts, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, after=6, cell=None):
    p = (cell or doc).add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(0)
    for t, bold in parts:
        r = p.add_run(t)
        r.font.name, r.font.size, r.bold = "Times New Roman", Pt(size), bold
    return p


def main():
    cites = cite_labels()
    title, paras, figs = parse(TEX.read_text(encoding="utf-8"))
    fignum = {lab: i for i, (_, _, lab) in enumerate(figs, start=1)}
    missing = [f for f, _, _ in figs if not f.exists()]
    if missing or not PDF.exists():
        sys.exit("missing figures or PDF — run make_qdlca_figures.py and pdflatex first")

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)

    add(doc, [(detex(title, cites, fignum), True)], align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    words = 0
    for p in paras:
        parts = runs(p, cites, fignum)
        words += sum(len(re.findall(r"[A-Za-z0-9][^\s]*", t)) for t, _ in parts)
        add(doc, parts, after=6)
    print(f"body: {words} words (limit 400; title, captions, references excluded)")

    doc.add_page_break()
    nrows = (len(figs) + 1) // 2
    table = doc.add_table(rows=nrows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (png, cap, lab) in enumerate(figs):
        cell = table.cell(i // 2, i % 2)
        width = Inches(3.05)
        if i == len(figs) - 1 and i % 2 == 0:      # odd count: last one centred
            cell = cell.merge(table.cell(i // 2, 1))
            width = Inches(3.3)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(png), width=width)
        add(doc, [(f"Figure {i + 1}. ", True), (detex(cap, cites, fignum), False)],
            size=9, after=8, cell=cell)

    add(doc, [("References", True)], size=11, after=4)
    for r in references_from_pdf():
        add(doc, [(r, False)], size=10, after=3)

    cp = doc.core_properties
    cp.author = cp.last_modified_by = cp.title = cp.subject = cp.comments = ""
    cp.category = cp.keywords = ""
    cp.revision = 1
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
