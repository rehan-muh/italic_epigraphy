#!/usr/bin/env python3
"""
Build abstract/qdlca27-abstract.docx — the file the call for papers asks for.

The text is converted from abstract/qdlca27-abstract.tex rather than retyped,
and the figures are the PNGs that scripts/make_abstract_figures.py renders from
the same drawing primitives as the TikZ ones, so the .docx cannot drift from
the PDF. Run the figure script first.

Submission is anonymous, so the document properties are blanked as well: Word
otherwise writes the machine's registered user into author/last-modified-by.

    py -3.12 scripts/make_abstract_docx.py
"""
from __future__ import annotations

import re
import sys
import unicodedata
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
TEX = ROOT / "abstract" / "qdlca27-abstract.tex"
OUT = ROOT / "abstract" / "qdlca27-abstract.docx"
FIGS = [ROOT / "abstract" / f"fig{i}.png" for i in (1, 2, 3)]
LIMIT = 400          # words, excluding title, captions and references

# The abstract's only display maths, spelled out once so the .docx is not at the
# mercy of a general LaTeX-to-Unicode converter.
MATH = {
    r"H(p)=-p\log_2p-(1-p)\log_2(1-p)": "H(p) = \u2212p log\u2082 p \u2212 (1\u2212p) log\u2082(1\u2212p)",
}


def demath(src: str) -> str:
    if src in MATH:
        return MATH[src]
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", src)
    s = s.replace(r"\log_2", "log\u2082 ").replace("-", "\u2212")
    return s.replace("\\", "")


# \`e and friends: the accent macro's combining character, composed afterwards.
ACCENT = {"`": "\u0300", "'": "\u0301", "^": "\u0302", '"': "\u0308", "~": "\u0303"}


def detex(s: str) -> str:
    """Plain-text substitutions; the markup macros are handled by runs()."""
    s = re.sub(r"\$(.+?)\$", lambda m: demath(m.group(1)), s)
    s = re.sub(r"\\([`'^\"~])\{?([a-zA-Z])\}?",
               lambda m: unicodedata.normalize("NFC", m.group(2) + ACCENT[m.group(1)]), s)
    s = re.sub(r"\\(?=\s)", " ", s)          # \  \u2014 a forced space, often line-broken
    s = s.replace("``", "\u201c").replace("''", "\u201d")
    s = s.replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace("{,}", ",").replace(r"\,", "\u2009").replace(r"\&", "&")
    s = s.replace(r"\%", "%").replace("~", "\u00a0")
    return re.sub(r"\s+", " ", s)            # collapse, but keep the run's edge spaces


def runs(s: str):
    """Split LaTeX into (text, italic, bold, monospace) runs."""
    out, pos = [], 0
    pat = re.compile(r"\\(textit|emph|textbf|textsc|texttt)\{([^{}]*)\}")
    for m in pat.finditer(s):
        if m.start() > pos:
            out.append((detex(s[pos:m.start()]), False, False, False))
        kind, body = m.group(1), detex(m.group(2))
        out.append((body.upper() if kind == "textsc" else body,
                    kind in ("textit", "emph"), kind == "textbf", kind == "texttt"))
        pos = m.end()
    if pos < len(s):
        out.append((detex(s[pos:]), False, False, False))
    out = [r for r in out if r[0]]
    if out:                                   # trim only the paragraph's own edges
        out[0] = (out[0][0].lstrip(),) + out[0][1:]
        out[-1] = (out[-1][0].rstrip(),) + out[-1][1:]
    return [r for r in out if r[0]]


def parse(tex: str):
    """Pull the title, body paragraphs, figure captions and references out."""
    src = "\n".join(l for l in tex.splitlines() if not l.lstrip().startswith("%"))

    title = re.search(r"\\title\{(.*?)\}\s*\n\\date", src, re.S).group(1)
    title = re.sub(r"\\vspace\{[^}]*\}|\\bfseries", "", title).strip()

    body_src = src.split(r"\begin{document}", 1)[1]
    body_src = body_src.split(r"\begin{figure}", 1)[0]
    body_src = re.sub(r"\\maketitle|\\vspace\{[^}]*\}", "", body_src)
    body = [p.strip() for p in body_src.split("\n\n") if p.strip()]

    captions = [re.sub(r"\s+", " ", c).strip()
                for c in re.findall(r"\\caption\{(.*?)\}\s*\n\\end\{figure\}", src, re.S)]

    refs_src = src.split(r"\textbf{References.}", 1)[1].split(r"\end{document}")[0]
    refs = [p.strip() for p in refs_src.strip().split("\n\n") if p.strip()]
    return title, body, captions, refs


def add(doc, text, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8,
        bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(space_after), Pt(0)
    for t, it, bd, mono in runs(text) if isinstance(text, str) else text:
        r = p.add_run(t)
        r.font.name = "Courier New" if mono else "Times New Roman"
        r.font.size = Pt(size - 1 if mono else size)
        r.italic, r.bold = it or italic, bd or bold
        if color:
            r.font.color.rgb = color
    return p


def main():
    missing = [f for f in FIGS if not f.exists()]
    if missing:
        sys.exit(f"missing {', '.join(f.name for f in missing)} — "
                 f"run: py -3.12 scripts/make_abstract_figures.py")

    title, body, captions, refs = parse(TEX.read_text(encoding="utf-8"))

    words = sum(len(detex(re.sub(r"\\[a-z]+\{([^{}]*)\}", r"\1", p)).split())
                for p in body)
    print(f"body: {words} words (limit {LIMIT}, title/captions/references excluded)")
    if words > LIMIT:
        print(f"  OVER by {words - LIMIT}")

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.4)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(11)

    add(doc, title, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, bold=True)
    for p in body:
        add(doc, p, space_after=8)

    for i, (fig, cap) in enumerate(zip(FIGS, captions), start=1):
        from PIL import Image
        with Image.open(fig) as im:
            w_cm = im.width / im.info.get("dpi", (300, 300))[0] * 2.54
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(6), Pt(2)
        p.add_run().add_picture(str(fig), width=Cm(min(w_cm, 16.2)))
        add(doc, [("Figure %d. " % i, False, True, False)] + runs(cap),
            size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)

    add(doc, r"\textbf{References.}", size=9, space_after=2)
    for r in refs:
        add(doc, r, size=9, space_after=2)

    # anonymise: no author, no machine name, no title metadata
    cp = doc.core_properties
    cp.author = cp.last_modified_by = cp.title = cp.subject = cp.comments = ""
    cp.category = cp.keywords = ""
    cp.revision = 1

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
